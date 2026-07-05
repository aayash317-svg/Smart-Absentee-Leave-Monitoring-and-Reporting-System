from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any

class ExcelService:
    @staticmethod
    def generate_students_sheet(students: List[Dict[str, Any]]) -> bytes:
        """Export a list of students and their attendance metrics to a formatted Excel file."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Directory"
        
        # Design system styles (matching the ERP color themes)
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Segoe UI", size=10)
        
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        headers = [
            "Register No", 
            "Enroll No", 
            "Full Name", 
            "Program", 
            "Department", 
            "Semester", 
            "Attendance %", 
            "Total Fine (INR)"
        ]
        
        ws.append(headers)
        
        # Format Headers
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        # Append Rows
        for r_idx, s in enumerate(students, start=2):
            row = [
                s.get("registerNo", ""),
                s.get("enrollNo", ""),
                s.get("fullName", ""),
                s.get("program", ""),
                s.get("departmentName", ""),
                s.get("semester", ""),
                s.get("attendancePercentage", 100.0),
                s.get("totalFine", 0.0)
            ]
            ws.append(row)
            
            # Format Data Row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if col_idx in [1, 2, 6]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        # Apply Column Autowidth
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        buffer = BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes
